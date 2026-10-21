"""Builds the Excel tracker. Conversions and summaries are live formulas driven by the FX sheet."""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from tracker import FX, INSTRUMENTS, ISSUER_TYPES, STATUSES

OUT = Path(__file__).parent / "output" / "India_Green_Finance_Tracker.xlsx"
FONT = "Arial"
H = Font(name=FONT, bold=True, color="FFFFFF", size=10)
HFILL = PatternFill("solid", fgColor="2E6B4F")
B, BB = Font(name=FONT, size=10), Font(name=FONT, size=10, bold=True)
BLUE = Font(name=FONT, size=10, color="0000FF")
TITLE = Font(name=FONT, size=14, bold=True, color="2E6B4F")
SUB = Font(name=FONT, size=10, italic=True, color="555555")
LINK = Font(name=FONT, size=10, color="0563C1", underline="single")
YEL = PatternFill("solid", fgColor="FFF9DB")
side = Side(style="thin", color="D0D7DE")
BOX = Border(left=side, right=side, top=side, bottom=side)
WRAP = Alignment(wrap_text=True, vertical="top")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
NUM = "#,##0;(#,##0);-"


def hdr(ws, row, labels, widths):
    for c, (lab, w) in enumerate(zip(labels, widths), 1):
        cell = ws.cell(row=row, column=c, value=lab)
        cell.font, cell.fill, cell.alignment, cell.border = H, HFILL, CEN, BOX
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[row].height = 30


def put(ws, r, c, v, font=B, fmt=None, align=WRAP):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font, cell.border, cell.alignment = font, BOX, align
    if fmt:
        cell.number_format = fmt
    return cell


def build(df: pd.DataFrame) -> Path:
    wb = Workbook()

    # ---------------------------------------------------------------- Dashboard (filled last)
    dash = wb.active
    dash.title = "Dashboard"

    # ---------------------------------------------------------------- Deals
    dl = wb.create_sheet("Deals")
    cols = ["Deal ID", "Date", "Issuer / borrower", "Issuer type", "Sector", "Instrument", "Green label",
            "Currency", "Amount (mn, deal ccy)", "USD mn", "INR crore", "Year", "Tenor (yrs)", "Pricing",
            "Lenders / investors", "Use of proceeds / KPIs", "Status", "Source"]
    hdr(dl, 1, cols, [11, 11, 26, 12, 18, 18, 11, 9, 13, 11, 11, 7, 9, 22, 30, 50, 11, 12])
    df = df.sort_values("date").reset_index(drop=True)
    n = len(df)
    for i, d in df.iterrows():
        r = i + 2
        put(dl, r, 1, d.deal_id, BB)
        put(dl, r, 2, d.date.to_pydatetime(), BLUE, "dd-mmm-yy", CEN)
        for c, k in ((3, "issuer"), (4, "issuer_type"), (5, "sector"), (6, "instrument"), (7, "green_label"), (8, "currency")):
            put(dl, r, c, d[k], BLUE, align=CEN if c in (4, 7, 8) else WRAP)
        put(dl, r, 9, None if pd.isna(d.amount_mn) else float(d.amount_mn), BLUE, NUM, CEN)
        put(dl, r, 10, f'=IF(I{r}="","",I{r}/INDEX(FX!$B$4:$B$7,MATCH(H{r},FX!$A$4:$A$7,0)))', B, NUM, CEN)
        put(dl, r, 11, f'=IF(J{r}="","",J{r}*FX!$B$5/10)', B, NUM, CEN)
        put(dl, r, 12, f"=YEAR(B{r})", B, "0", CEN)
        put(dl, r, 13, "" if pd.isna(d.tenor_years) else str(d.tenor_years), BLUE, align=CEN)
        for c, k in ((14, "pricing"), (15, "lenders_investors"), (16, "use_of_proceeds_or_kpis"), (17, "status")):
            put(dl, r, c, "" if pd.isna(d[k]) else d[k], BLUE, align=CEN if c == 17 else WRAP)
        s = put(dl, r, 18, "Link", LINK)
        s.hyperlink = d.source_url
        dl.row_dimensions[r].height = 48
        for c in range(1, 19):
            if dl.cell(row=r, column=c).font.color and dl.cell(row=r, column=c).font.color.rgb == "000000FF":
                dl.cell(row=r, column=c).fill = YEL
    last = n + 1
    # empty input rows for new deals, formulas pre-filled
    for r in range(last + 1, last + 16):
        for c in range(1, 19):
            put(dl, r, c, None, BLUE).fill = YEL
        dl.cell(row=r, column=10, value=f'=IF(I{r}="","",I{r}/INDEX(FX!$B$4:$B$7,MATCH(H{r},FX!$A$4:$A$7,0)))').number_format = NUM
        dl.cell(row=r, column=11, value=f'=IF(J{r}="","",J{r}*FX!$B$5/10)').number_format = NUM
        dl.cell(row=r, column=12, value=f'=IF(B{r}="","",YEAR(B{r}))')
        for c in (10, 11, 12):
            dl.cell(row=r, column=c).font, dl.cell(row=r, column=c).fill = B, PatternFill()
    END = last + 15
    for col, opts in (("D", ISSUER_TYPES), ("F", INSTRUMENTS), ("G", ["Confirmed", "Unconfirmed"]),
                      ("H", list(FX)), ("Q", STATUSES)):
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
        dv.add(f"{col}2:{col}{END}")
        dl.add_data_validation(dv)
    dl.freeze_panes = "D2"
    dl.auto_filter.ref = f"A1:R{END}"

    # ---------------------------------------------------------------- FX
    fx = wb.create_sheet("FX")
    fx["A1"], fx["A1"].font = "FX assumptions (units per 1 USD)", TITLE
    fx["A2"], fx["A2"].font = "Approximate October 2026 levels. Blue cells are inputs: edit to re-price every deal.", SUB
    hdr(fx, 3, ["Currency", "Per USD"], [14, 12])
    for i, (ccy, rate) in enumerate(FX.items(), 4):
        put(fx, i, 1, ccy, BB)
        put(fx, i, 2, rate, BLUE, "#,##0.00", CEN).fill = YEL
    # INR must stay in row 5 (INR crore formulas reference FX!$B$5)
    assert list(FX)[1] == "INR"

    # ---------------------------------------------------------------- Dashboard
    D = f"Deals!$J$2:$J${END}"
    ws = dash
    ws["A1"], ws["A1"].font = "India Green Bond & ESG Loan Tracker", TITLE
    ws["A2"], ws["A2"].font = ("Labelled green, ESG and sustainability-linked debt by Indian issuers, Jan 2023 - Oct 2026. "
                               "Totals count completed deals (Priced, Signed, Devolved) in USD mn at FX-sheet rates."), SUB
    ws.column_dimensions["A"].width = 30
    done = lambda extra: (f'SUMIFS({D},Deals!$Q$2:$Q${END},"Priced"{extra})+SUMIFS({D},Deals!$Q$2:$Q${END},"Signed"{extra})'
                          f'+SUMIFS({D},Deals!$Q$2:$Q${END},"Devolved"{extra})')
    cnt = lambda extra: (f'COUNTIFS(Deals!$Q$2:$Q${END},"Priced"{extra})+COUNTIFS(Deals!$Q$2:$Q${END},"Signed"{extra})'
                         f'+COUNTIFS(Deals!$Q$2:$Q${END},"Devolved"{extra})')

    kpis = [("Completed deals", f"={cnt('')}", "0"),
            ("Total raised (USD mn)", f"={done('')}", NUM),
            ("Total raised (INR crore)", f"=B5*FX!$B$5/10", NUM),
            ("Largest completed deal (USD mn)", f'=MAX(_xlfn.MAXIFS({D},Deals!$Q$2:$Q${END},"Priced"),_xlfn.MAXIFS({D},Deals!$Q$2:$Q${END},"Signed"),_xlfn.MAXIFS({D},Deals!$Q$2:$Q${END},"Devolved"))', NUM),
            ("Sovereign green bonds outstanding (INR bn, Aug 2026)", 877, NUM)]
    for i, (lab, f, fmt) in enumerate(kpis, 4):
        put(ws, i, 1, lab, BB)
        c = put(ws, i, 2, f, BLUE if isinstance(f, (int, float)) else BB, fmt, CEN)
    ws.cell(row=8, column=3, value="Source: Reuters via Business Recorder, 14 Aug 2026").font = SUB
    ws.column_dimensions["B"].width = 16

    def table(top, title, key_col, keys):
        put(ws, top, 1, title, H).fill = HFILL
        for c, lab in ((2, "Deals"), (3, "USD mn"), (4, "Share")):
            put(ws, top, c, lab, H, align=CEN).fill = HFILL
        for i, k in enumerate(keys, top + 1):
            crit = f',Deals!${key_col}$2:${key_col}${END},$A{i}'
            put(ws, i, 1, k)
            put(ws, i, 2, f"={cnt(crit)}", B, "0", CEN)
            put(ws, i, 3, f"={done(crit)}", B, NUM, CEN)
            put(ws, i, 4, f"=IF($B$5=0,0,C{i}/$B$5)", B, "0%", CEN)
        tot = top + len(keys) + 1
        put(ws, tot, 1, "Total", BB)
        put(ws, tot, 2, f"=SUM(B{top + 1}:B{tot - 1})", BB, "0", CEN)
        put(ws, tot, 3, f"=SUM(C{top + 1}:C{tot - 1})", BB, NUM, CEN)
        put(ws, tot, 4, f"=SUM(D{top + 1}:D{tot - 1})", BB, "0%", CEN)
        return top, tot

    for c in "CD":
        ws.column_dimensions[c].width = 12
    t1 = table(10, "By instrument", "F", INSTRUMENTS)
    t2 = table(t1[1] + 2, "By issuer type", "D", ISSUER_TYPES)
    years = [2023, 2024, 2025, 2026]
    t3 = table(t2[1] + 2, "By year", "L", years)

    for top, tot, title, anchor in ((t1[0], t1[1], "USD mn by instrument", "F4"),
                                    (t2[0], t2[1], "USD mn by issuer type", "F21"),
                                    (t3[0], t3[1], "USD mn by year", "F38")):
        ch = BarChart()
        ch.type, ch.title, ch.legend = "bar", title, None
        ch.add_data(Reference(ws, min_col=3, min_row=top, max_row=tot - 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=top + 1, max_row=tot - 1))
        ch.height, ch.width = 7.5, 14
        ws.add_chart(ch, anchor)

    # ---------------------------------------------------------------- Sovereign
    sv = wb.create_sheet("Sovereign Greenium")
    sv["A1"], sv["A1"].font = "Sovereign green bonds: auction outcomes and greenium", TITLE
    sv["A2"], sv["A2"].font = ("Greenium = yield discount vs a comparable conventional G-sec, in basis points. "
                               "Blank where the source gives none."), SUB
    hdr(sv, 4, ["Auction", "Tenor", "Notified (INR cr)", "Outcome", "Greenium (bps)", "What it tells you", "Source"],
        [14, 9, 14, 30, 12, 60, 10])
    rows = [
        ("Jan 2023", "5y / 10y", 8000, "Debut; >4x subscribed", 6,
         "Strong debut demand, priced through conventional curve", "https://renewablewatch.in/2025/07/01/green-yields/"),
        ("May 2024", "", 6000, "Cancelled: no pricing advantage", None,
         "Government refused to issue without a greenium", "https://renewablewatch.in/2025/07/01/green-yields/"),
        ("Aug 2024", "", 6000, "Only Rs 1,697 cr bid; cut-off 6.90%", None,
         "Thin demand at short/medium tenors", "https://renewablewatch.in/2025/07/01/green-yields/"),
        ("Nov 2024", "", 5000, "70% devolved on primary dealers (6.78-6.84%)", None,
         "Banks unwilling to pay a greenium", "https://renewablewatch.in/2025/07/01/green-yields/"),
        ("Jan 2025", "10y", None, "Rs 3,945 cr devolved on primary dealers", None,
         "10-year green paper struggles", "https://www.business-standard.com/finance/news/rbi-partially-devolves-rs-3-945-cr-worth-of-10-year-sovereign-green-bonds-125013101508_1.html"),
        ("Jun 2025", "30y", 5000, "All bids rejected", None,
         "Investors sought higher yields than the government would pay", "https://www.business-standard.com/economy/news/rbi-rejects-30-year-green-bond-bids-as-investors-seek-higher-yields-125061300999_1.html"),
        ("Apr 2026", "30y", None, "~2.5x subscribed; 7.50% coupon", 6,
         "Largest greenium since 2023; insurers drive demand via 'infra' classification", "https://www.brecorder.com/news/amp/40418096"),
        ("Aug 2026", "30y", 5000, "Sold in full", 4,
         "H1 FY27 average greenium 4 bps, highest since launch", "https://www.brecorder.com/news/amp/40434824"),
        ("Oct 2026 - Jan 2027", "30y", 15000, "Scheduled: 5 x Rs 3,000 cr", None,
         "Government shifted green supply to 30-year, where demand exists", "https://goldenpi.com/blog/bond-news/india-sovereign-green-bonds-h2-fy27-15000-crore-issuance-schedule-30-year-auctions/"),
    ]
    for i, row in enumerate(rows, 5):
        for c, v in enumerate(row[:6], 1):
            put(sv, i, c, v, BLUE if c in (3, 5) else B, NUM if c == 3 else ("0" if c == 5 else None),
                CEN if c in (2, 3, 5) else WRAP)
        put(sv, i, 7, "Link", LINK).hyperlink = row[6]
        sv.row_dimensions[i].height = 30
    e = 4 + len(rows)
    put(sv, e + 1, 4, "Average greenium where reported (bps)", BB)
    put(sv, e + 1, 5, f"=AVERAGE(E5:E{e})", BB, "0.0", CEN)

    # ---------------------------------------------------------------- About
    ab = wb.create_sheet("About")
    ab.column_dimensions["A"].width = 110
    notes = [
        ("About this tracker", TITLE),
        ("Scope: labelled green, ESG and sustainability-linked bonds and loans by Indian issuers and borrowers, plus "
         "sovereign green bond auctions, Jan 2023 to Oct 2026. Every deal row links to its public source.", B),
        ("Inputs: blue text on yellow is editable. Add a deal in the empty yellow rows at the bottom of the Deals sheet; "
         "USD, INR crore and every Dashboard figure update automatically.", B),
        ("Completed = Priced, Signed or Devolved. Rejected and Scheduled deals are listed but not counted in totals.", B),
        ("Green label 'Unconfirmed' = the issuer is a green-sector lender and a trade publication calls the issue green, "
         "but the issuer's own release does not. Treat with caution.", B),
        ("FX: deals are converted at the FX sheet's rates, not the historical rate on each deal date, so totals are "
         "indicative.", B),
        ("Limits: this is a curated sample from public news and filings, not a complete market database "
         "(e.g. Bloomberg, LSEG, CBI). Amounts not disclosed by the source are left blank and excluded from totals.", B),
        ("", B),
        ("Prepared by Mayank Majoka, October 2026. Not investment advice.", SUB),
    ]
    for i, (t, f) in enumerate(notes, 1):
        c = ab.cell(row=i, column=1, value=t)
        c.font, c.alignment = f, WRAP

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    return OUT
